"""INEP: resultados do IDEB (Brasil), divulgação mais recente.

O servidor download.inep.gov.br derruba a conexão TLS do Python, então o
download é feito com curl.
"""
from __future__ import annotations

import io
import re
import zipfile

import openpyxl
import pandas as pd

from etl.comum import Bruto, obter

EDICAO = 2025
URL = "https://download.inep.gov.br/ideb/resultados/divulgacao_brasil_ideb_{ed}.zip"
URL_PAGINA = "https://www.gov.br/inep/pt-br/areas-de-atuacao/pesquisas-estatisticas-e-indicadores/ideb/resultados"

ETAPAS = {
    "Anos Iniciais": "anos_iniciais",
    "Anos Finais": "anos_finais",
    "(EM)": "ensino_medio",
}


def ideb_brasil(edicao: int = EDICAO) -> tuple[pd.DataFrame, Bruto]:
    """IDEB observado, rede Total, por etapa e ano de aplicação."""
    bruto = obter("inep", f"divulgacao_brasil_ideb_{edicao}.zip", URL.format(ed=edicao), usar_curl=True, timeout=600)
    z = zipfile.ZipFile(io.BytesIO(bruto.conteudo))
    nome = next(n for n in z.namelist() if n.endswith(".xlsx"))
    wb = openpyxl.load_workbook(io.BytesIO(z.read(nome)), read_only=True, data_only=True)
    linhas = []
    for aba in wb.sheetnames:
        etapa = next((v for k, v in ETAPAS.items() if k in aba), None)
        if etapa is None:
            continue
        rows = list(wb[aba].iter_rows(values_only=True))
        # Usa o cabeçalho legível ("IDEB\n2023\n(N x P)"): no arquivo de 2025 o
        # cabeçalho técnico VL_OBSERVADO_2023 dos anos finais está em branco.
        cab = next(r for r in rows if r and r[0] == "Brasil" and r[1] == "Rede")
        total = next(r for r in rows if r and r[0] == "Brasil" and str(r[1]).strip() == "Total")
        for i, c in enumerate(cab):
            m = re.match(r"IDEB\s*(\d{4})", str(c or ""))
            if m and total[i] not in (None, "-", ""):
                linhas.append({"etapa": etapa, "ano": int(m.group(1)), "valor": float(str(total[i]).replace(",", "."))})
    return pd.DataFrame(linhas), bruto
