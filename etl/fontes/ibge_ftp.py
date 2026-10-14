"""IBGE (FTP): Síntese de Indicadores Sociais (pobreza) e Projeção da População."""
from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pandas as pd

from etl.comum import Bruto, obter

FTP = "https://ftp.ibge.gov.br/"
SIS_EDICAO = 2025
URL_SIS = FTP + "Indicadores_Sociais/Sintese_de_Indicadores_Sociais/Sintese_de_Indicadores_Sociais_{ed}/Tabelas/ods/2_Distribuicao_Renda_ods.zip"
URL_SIS_PAGINA = "https://www.ibge.gov.br/estatisticas/sociais/populacao/9221-sintese-de-indicadores-sociais.html"
URL_PROJECAO = FTP + "Projecao_da_Populacao/Projecao_da_Populacao_2024/projecoes_2024_tab4_indicadores.xlsx"

T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
P = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p"


def _ods_planilhas(conteudo: bytes) -> dict[str, list[list[str]]]:
    raiz = ET.fromstring(zipfile.ZipFile(io.BytesIO(conteudo)).read("content.xml"))
    out = {}
    for sh in raiz.iter(T + "table"):
        linhas = []
        for r in sh.iter(T + "table-row"):
            cel = []
            for c in r.iter(T + "table-cell"):
                rep = min(int(c.get(T + "number-columns-repeated", "1")), 50)
                cel += [" ".join("".join(p.itertext()) for p in c.findall(P))] * rep
            linhas.append(cel)
        out[sh.get(T + "name")] = linhas
    return out


def _num_br(s: str) -> float:
    s = (s or "").strip().replace(" ", "").replace("\xa0", "")
    if s in ("", "-", "..", "...", "x"):
        return float("nan")
    return float(s.replace(".", "").replace(",", "."))


def sis_pobreza(edicao: int = SIS_EDICAO) -> tuple[pd.DataFrame, Bruto]:
    """Tabela 2.17: proporção de pessoas abaixo das linhas do Banco Mundial
    (US$ 2,15 e US$ 6,85 PPC 2017), Brasil, uma aba por ano."""
    bruto = obter("ibge_sis", f"SIS{edicao}_2_Distribuicao_Renda_ods.zip", URL_SIS.format(ed=edicao), timeout=600)
    z = zipfile.ZipFile(io.BytesIO(bruto.conteudo))
    nome = next(n for n in z.namelist() if re.search(r"Tabela 2\.17 ", n))
    abas = _ods_planilhas(z.read(nome))
    linhas = []
    for aba, rows in abas.items():
        if not re.fullmatch(r"\d{4}", aba.strip()):
            continue  # ignora abas de coeficiente de variação
        cab = [c for c in next(r for r in rows if any("US$ 2,15" in c for c in r)) if c.strip()]
        i215 = next(i for i, c in enumerate(cab) if "US$ 2,15" in c)
        i685 = next(i for i, c in enumerate(cab) if "US$ 6,85" in c)
        tot = next(r for r in rows if r and r[0].strip() == "Total")
        # Linha "Total": rótulo, população (mil) e as proporções na mesma
        # ordem do cabeçalho das classes de rendimento.
        vals = [c for c in tot[1:] if c.strip()]
        linhas.append({
            "ano": int(aba),
            "populacao_mil": _num_br(vals[0]),
            "extrema_pobreza": _num_br(vals[1 + i215]),
            "pobreza": _num_br(vals[1 + i685]),
        })
    return pd.DataFrame(linhas).sort_values("ano").reset_index(drop=True), bruto


def projecao_populacao_brasil() -> tuple[pd.DataFrame, Bruto]:
    """População total e feminina do Brasil por ano (Projeção da População, revisão 2024)."""
    bruto = obter("ibge_projecao", "projecoes_2024_tab4_indicadores.xlsx", URL_PROJECAO, timeout=600)
    wb = openpyxl.load_workbook(io.BytesIO(bruto.conteudo), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    cab, linhas = None, []
    for row in ws.iter_rows(values_only=True):
        if cab is None:
            if row and row[0] == "ANO":
                cab = list(row)
            continue
        if row and row[cab.index("SIGLA")] == "BR":
            linhas.append({"ano": int(row[0]), "populacao": float(row[cab.index("POP_T")]), "populacao_mulheres": float(row[cab.index("POP_M")])})
    return pd.DataFrame(linhas), bruto
