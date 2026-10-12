"""ANP: Levantamento de Preços de Combustíveis, série histórica mensal (Brasil)."""
from __future__ import annotations

import io

import openpyxl
import pandas as pd

from etl.comum import Bruto, obter

URL_MENSAL_BRASIL = (
    "https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/"
    "precos-revenda-e-de-distribuicao-combustiveis/shlp/mensal/mensal-brasil-desde-jan2013.xlsx"
)
URL_PAGINA = (
    "https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/"
    "precos-revenda-e-de-distribuicao-combustiveis/serie-historica-do-levantamento-de-precos"
)


def mensal_brasil() -> tuple[pd.DataFrame, Bruto]:
    bruto = obter("anp", "mensal-brasil-desde-jan2013.xlsx", URL_MENSAL_BRASIL, timeout=300)
    wb = openpyxl.load_workbook(io.BytesIO(bruto.conteudo), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas, cab = [], None
    for row in ws.iter_rows(values_only=True):
        if cab is None:
            if row and row[0] == "MÊS":
                cab = [str(c).strip() if c else "" for c in row]
            continue
        if not row or row[0] is None:
            continue
        linhas.append(dict(zip(cab, row)))
    df = pd.DataFrame(linhas)
    df = df.rename(columns={"MÊS": "mes", "PRODUTO": "produto", "UNIDADE DE MEDIDA": "unidade", "PREÇO MÉDIO REVENDA": "preco"})
    df["mes"] = pd.to_datetime(df["mes"])
    df["preco"] = pd.to_numeric(df["preco"], errors="coerce")
    return df[["mes", "produto", "unidade", "preco"]], bruto
